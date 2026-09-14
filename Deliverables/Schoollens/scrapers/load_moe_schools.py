"""Load MOE_Approve_List.xlsx into schools + moe raw_sources (Section 10 / 16)."""

import hashlib
import json
import os
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from supabase import create_client

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(
    r"D:\SMT\Personal\myPKA-main\myPKA-main\Owner Inbox\App\School Data\MOE_Approve_List.xlsx"
)


def _load_env():
    for env_path in (ROOT / "web" / ".env.local", ROOT / "rag-service" / ".env", ROOT / "rag-service" / ".env.example"):
        if not env_path.exists():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))
    if not os.environ.get("SUPABASE_URL"):
        os.environ["SUPABASE_URL"] = os.environ.get("NEXT_PUBLIC_SUPABASE_URL", "")


def _client():
    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        raise SystemExit("SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY must be set")
    return create_client(url, key)


def _compact(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _clean(value):
    text = str(value or "").replace("\xa0", " ").strip()
    if text in {"", "0", "none", "None"}:
        return None
    return text


def _clean_url(value, facebook=False):
    text = _clean(value)
    if not text:
        return None
    lower = text.lower()
    if not facebook and ("facebook.com" in lower or "fb.com" in lower):
        return None
    if facebook and "facebook.com" not in lower and "fb.com" not in lower:
        return None
    return text


def _za_digits(text):
    table = str.maketrans("၀၁၂၃၄၅၆၇၈၉", "0123456789")
    return str(text or "").translate(table)


def _approval_dates(period):
    years = re.findall(r"20\d{2}", _za_digits(period))
    if len(years) >= 2:
        return f"{years[0]}-06-01", f"{years[-1]}-05-31"
    if years:
        return f"{years[0]}-06-01", None
    return None, None


def _brand_key(name, website):
    if website:
        host = re.sub(r"^https?://(www\.)?", "", website.lower()).split("/")[0]
        if host and "facebook." not in host:
            return f"site:{host}"
    compact = _compact(re.sub(r"\([^)]*\)", " ", name or ""))
    compact = re.sub(r"\d+$", "", compact)
    aliases = (
        ("ilbcinternationalschool", "ilbc"),
        ("ilbc", "ilbc"),
        ("higherchampsinternationalschoolhcis", "hcis"),
        ("higherchampsinternationalschool", "hcis"),
        ("hcis", "hcis"),
    )
    for prefix, key in aliases:
        if compact.startswith(prefix) or compact == prefix:
            return f"brand:{key}"
    return f"brand:{compact}" if compact else None


def _group_id(key):
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"schoollens:{key}"))


def _read_xlsx(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        first = str(row[0] or "").strip()
        if not first.isdigit():
            continue
        name = _clean(row[1])
        if not name:
            continue
        website = _clean_url(row[4], facebook=False)
        facebook = _clean_url(row[7], facebook=True) or _clean_url(row[4], facebook=True)
        start, end = _approval_dates(row[3])
        rows.append(
            {
                "seq": first,
                "name": name,
                "address": _clean(row[2]),
                "period": _clean(row[3]),
                "official_website_url": website,
                "official_facebook_url": facebook,
                "moe_approved_from": start,
                "moe_approved_to": end,
                "curriculum_type": "international",
            }
        )
    return rows


def _match_existing(existing, listing):
    name_key = _compact(listing["name"])
    addr_key = _compact(listing.get("address"))
    for school in existing:
        if _compact(school.get("name")) == name_key and _compact(school.get("address")) == addr_key:
            return school
    if addr_key:
        hits = [school for school in existing if _compact(school.get("address")) == addr_key]
        if len(hits) == 1:
            return hits[0]
    return None


def main():
    _load_env()
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not path.exists():
        raise SystemExit(f"MOE list not found: {path}")
    listings = _read_xlsx(path)
    if not listings:
        raise SystemExit("no school rows in the workbook")
    db = _client()
    existing = db.table("schools").select("id, name, address").execute().data or []
    used_ids = set()
    now = datetime.now(timezone.utc).isoformat()
    inserted = 0
    updated = 0
    moe_rows = 0
    for listing in listings:
        key = _brand_key(listing["name"], listing["official_website_url"])
        group_id = _group_id(key) if key else None
        payload = {
            "name": listing["name"],
            "address": listing["address"],
            "moe_approved_from": listing["moe_approved_from"],
            "moe_approved_to": listing["moe_approved_to"],
            "official_website_url": listing["official_website_url"],
            "official_facebook_url": listing["official_facebook_url"],
            "curriculum_type": listing["curriculum_type"],
            "school_group_id": group_id,
        }
        match = _match_existing([row for row in existing if row["id"] not in used_ids], listing)
        if match:
            db.table("schools").update(payload).eq("id", match["id"]).execute()
            school_id = match["id"]
            used_ids.add(school_id)
            updated += 1
        else:
            saved = db.table("schools").insert(payload).execute()
            school_id = saved.data[0]["id"]
            existing.append({"id": school_id, "name": listing["name"], "address": listing["address"]})
            used_ids.add(school_id)
            inserted += 1

        raw = {
            "school_id": school_id,
            "source_type": "moe",
            "seq": listing["seq"],
            "period": listing["period"],
            "extracted_text": (
                f"School: {listing['name']}\n"
                f"Address: {listing['address'] or ''}\n"
                f"MOE approval period: {listing['period'] or ''}\n"
                f"Official website: {listing['official_website_url'] or 'none'}\n"
                f"Official Facebook: {listing['official_facebook_url'] or 'none'}"
            ),
            "crawled_at": now,
        }
        digest = "sha256:" + hashlib.sha256(json.dumps(raw, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        already = (
            db.table("raw_sources")
            .select("id")
            .eq("school_id", school_id)
            .eq("source_type", "moe")
            .eq("content_hash", digest)
            .limit(1)
            .execute()
            .data
        )
        if already:
            continue
        db.table("raw_sources").insert(
            {
                "school_id": school_id,
                "source_type": "moe",
                "crawl_status": "success",
                "raw_json": raw,
                "content_hash": digest,
                "crawled_at": now,
            }
        ).execute()
        moe_rows += 1

    leftover = [row for row in existing if row["id"] not in used_ids and "ilbc" in _compact(row.get("name"))]
    if leftover:
        ilbc_group = _group_id("site:ilbc.edu.mm")
        for row in leftover:
            db.table("schools").update({"school_group_id": ilbc_group, "curriculum_type": "international"}).eq(
                "id", row["id"]
            ).execute()
        print(f"linked leftover ILBC rows {len(leftover)}")

    print(f"listings {len(listings)} inserted {inserted} updated {updated} moe_raw_sources {moe_rows}")


if __name__ == "__main__":
    main()
